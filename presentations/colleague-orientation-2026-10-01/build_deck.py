from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
import json, html

OUT = Path(__file__).parent
R = Presentation(); R.slide_width=Inches(13.333); R.slide_height=Inches(7.5)
BG='F5F3EE'; NAVY='152D3A'; TEAL='166B70'; GOLD='C69944'; GRAY='556671'; PALE='E6EEEB'; WHITE='FFFFFF'; RED='AD5346'
slides=[]
def rect(s,x,y,w,h,color,line=None,radius=False):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    q.fill.solid(); q.fill.fore_color.rgb=RGBColor.from_string(color)
    q.line.fill.background() if not line else None
    if line:q.line.color.rgb=RGBColor.from_string(line)
    return q
def txt(s,x,y,w,h,text,size=20,color=NAVY,bold=False):
    q=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=q.text_frame
    tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(8)
    return q
def line(s,x1,y1,x2,y2,color=TEAL,width=2):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); q.line.color.rgb=RGBColor.from_string(color); q.line.width=Pt(width)
    return q
def arrow(s,x,y,w=.38,h=.25,color=TEAL):return rectshape(s,MSO_SHAPE.RIGHT_ARROW,x,y,w,h,color)
def rectshape(s,shape,x,y,w,h,color):
    q=s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(color);q.line.fill.background();return q
def new(title,kicker,sub,source,notes):
    s=R.slides.add_slide(R.slide_layouts[6]); s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    rect(s,.45,.38,.12,.28,GOLD);txt(s,.7,.36,11,.3,kicker.upper(),11,TEAL,True)
    txt(s,.65,.87,12,0.65,title,31,NAVY,True);txt(s,.67,1.62,12,.65,sub,17,GRAY)
    line(s,.65,6.91,12.65,6.91,'CDD4D2',1)
    txt(s,.67,7.05,11.4,.23,source,9,GRAY);txt(s,12,7.01,.7,.3,f'{len(slides)+1:02d} / 15',10,GRAY)
    s.notes_slide.notes_text_frame.text=notes
    slides.append({'title':title,'notes':notes,'source':source});return s
def card(s,x,y,w,h,title,body,accent=TEAL,sz=18):
    rect(s,x,y,w,h,WHITE,radius=True);rect(s,x,y,.06,h,accent)
    txt(s,x+.2,y+.18,w-.38,.5,title,20,accent,True)
    body_y=.61 if h<=1.3 else .82
    txt(s,x+.2,y+body_y,w-.4,max(.4,h-body_y-.1),body,sz)
def band(s,text,y=6.25):
    rect(s,.67,y,12,.43,PALE);txt(s,.85,y+.07,11.65,.3,text,14,TEAL,True)

s=new('Anchor: one place to carry work forward','Colleague orientation · Initial mockup v1','A personal AI workspace connecting goals, projects, conversations and specialist workflows.','ORIENTATION MOCKUP v1 · 1 OCTOBER 2026 · source-based, not a product certification',
'''Present this as John's working personal system. Anchor is Python-based software on PUTER, reached through browsers over Tailscale. The models run through their provider tools; the local host coordinates work and keeps project artifacts. This presentation explains the current source and recorded behavior, not a fresh full acceptance test. All UI panels in this deck are illustrative schematics with fictional example content, not screenshots. Sources: C:/dev/Anchor/CLAUDE.md (Runtime / Host); anchor_gui.py; high_seat_page.py; campaign_page.py. Suggested opening: “I can leave one device, open another, and return to the same workspace.”''')
txt(s,.8,2.6,6.5,2.15,'Keep the goal.\nKeep the context.\nMove the work.',35,NAVY,True)
for i,(a,b) in enumerate([('ORGANIZE','Goals and active initiatives'),('WORK','A cockpit for each effort'),('CHECK','Specialists with review built in')]):
    card(s,8,2.45+i*1.22,4.35,1.06,a,b,TEAL,16)
band(s,'A 15-slide guided tour, followed by a live demonstration.')

s=new('One host, several ways to reach it','01 · Access','PUTER does the coordinating. Your browser is the window into the same workspace.','SOURCES: Anchor/CLAUDE.md · anchor_gui.py · Tailscale documentation',
'''Host identity: PUTER and gwl-server refer to the same desktop in the runtime documentation. The persistent Windows service runs Python anchor_gui.py on local port 8777. Tailscale supplies private network access; “broadcasting” is an informal description, not public internet publication. Workstation, laptop, iPhone and iPad are the user's stated access devices. Device access depends on the host being available and the device/network being authorized. Remote model providers may process prompts: local hosting does not imply local-only AI inference. This deck did not retest each physical device. https://tailscale.com/kb/1151/what-is-tailscale''')
card(s,.8,2.55,3.05,2.95,'YOUR DEVICES','Work workstation\nLaptop\niPhone · iPad',TEAL,22)
arrow(s,4.1,3.7,.55,.34)
card(s,4.88,2.55,3.55,2.95,'PRIVATE CONNECTION','Tailscale\nAuthorized devices reach Anchor through a browser.',TEAL,21)
arrow(s,8.68,3.7,.55,.34)
card(s,9.45,2.55,3.05,2.95,'PUTER','Python web service\nProject files + records\nAI tools + skill engines',TEAL,19)
band(s,'One shared workspace across devices; AI providers are reached by the host’s configured tools.')

s=new('How the pieces fit together','02 · Organization','A goal says why. An initiative identifies the work. A cockpit is where you interact with it.','SOURCES: high_seat_page.py · campaign_page.py · campaign_parts.py · illustrative example',
'''The current product has overlapping views, not a rigid folder tree. The Cave of Wonders is the Aladdin persona's name for the High Seat portfolio home. The Campaign holds goals, priority tiers and component references. Components can reference a project cockpit, an effort cockpit within a project, a desk task, a project step, a deliverable or a date. Active and On deck are attention placements for work items, not separate copies of their project files and not automatic execution commands. Projects is the registry of project homes; an effort may have its own directory and cockpit. Example is fictional to keep the shared deck free of personal project details.''')
card(s,.8,2.45,3.3,1.25,'CAVE OF WONDERS','The portfolio home',TEAL,19)
card(s,4.65,2.45,3.3,1.25,'CAMPAIGN','Goals + priorities',TEAL,19);arrow(s,4.2,2.95)
card(s,8.5,2.45,3.95,1.25,'INITIATIVES','Work serving those goals',TEAL,19);arrow(s,8.05,2.95)
line(s,10.5,3.75,10.5,4.15)
card(s,.8,4.25,3.3,1.55,'PROJECT','Example: AI teaching',GOLD,19)
card(s,4.65,4.25,3.3,1.55,'EFFORT + COCKPIT','Example: build one workshop',GOLD,18);arrow(s,4.2,4.88)
card(s,8.5,4.25,3.95,1.55,'ACTIVE / ON DECK','Choose what deserves attention',GOLD,18)
band(s,'An initiative can open a whole project or one named effort inside it. Its goal ownership stays intact.')

s=new('The Cave of Wonders: your portfolio home','03 · The front door','See what needs attention, open the work, and reach the tools from one surface.','SOURCES: anchor_settings.py · high_seat_page.py · UI SCHEMATIC, not a screenshot',
'''Cave of Wonders is a selectable persona name, not a separate storage system. Ecgberht, Aladdin and Jarvis share one steward engine with different visual naming. Current source includes the Campaign goal tile, Active and On deck shelves, Projects registry, Desk, seeds, a brief/dialogue area, Workbench and access to Skill Foundry/Doctor. Calendar and Email are explicitly “not connected.” Do not present these as working integrations. Demo cue: start at the Cave, point to the goal summary, then expand a single active card. Source is the current renderer, rather than the stale July ANCHOR-STATE summary.''')
rect(s,.8,2.4,8,3.9,WHITE,radius=True)
txt(s,1.05,2.65,7.45,.4,'CAVE OF WONDERS',22,TEAL,True)
for x,y,w,h,t in [(1.05,3.25,7.45,.65,'Campaign · goals and priorities'),(1.05,4.08,3.6,.85,'Active · working now'),(4.9,4.08,3.6,.85,'On deck · attention next'),(1.05,5.12,3.6,.65,'Projects · all project homes'),(4.9,5.12,3.6,.65,'Desk · small tasks')]:
    rect(s,x,y,w,h,PALE,radius=True);txt(s,x+.14,y+.14,w-.28,h-.12,t,17,TEAL,True)
txt(s,9.15,2.6,3.2,3.95,'Orient\nRead the brief and current state.\n\nAct\nExpand work or talk to the steward.\n\nReach tools\nFoundry and Doctor.',18)

s=new('The Campaign turns ambitions into organized work','04 · Goals and major efforts','Example: “Help colleagues use AI well” becomes a goal with several concrete components.','SOURCES: campaign_page.py · high_seat_page.py · illustrative goal, not personal campaign data',
'''The Campaign view reads a common campaign record, with project references resolved from their own records. It provides three priority tiers: do first, next in line, when there is room. Goals can contain cockpit references, desk items, steps, deliverables and dates. Seeds hold early ideas; the Boneyard retains done or parked goals. The page has a steward conversation for discussing and changing organization. Seating an initiative on Active does not authorize starting its execution. Demo cue: open a goal and identify its completion condition, components and priority.''')
card(s,.8,2.4,7.3,1.05,'GOAL: HELP COLLEAGUES USE AI WELL','Done when: the team can complete one useful workflow.',TEAL,17)
for i,(a,b) in enumerate([('Workshop effort','Prepare examples and teaching materials'),('Orientation deck','Explain the workspace and its specialist tools'),('Desk task + date','Schedule the demonstration and invite participants')]):
    card(s,1.2,3.7+i*.78,6.9,.67,a,b,GOLD,13)
    # compact row treatment
    q=s.shapes[-2];q.text_frame.paragraphs[0].font.size=Pt(15)
    q.top=Inches(3.8+i*.78);q.height=Inches(.3);q.width=Inches(2.2)
    q=s.shapes[-1];q.top=Inches(3.8+i*.78);q.left=Inches(3.75);q.width=Inches(4.05);q.height=Inches(.4)
card(s,8.55,2.4,3.9,3.55,'PRIORITY TIERS','1  Do first\n2  Next in line\n3  When there is room\n\nKeep early ideas as seeds. Retain done or parked goals.',TEAL,20)

s=new('Expand an initiative and work right there','05 · Active and On deck','The card is a doorway into the same effort, conversation and deliverables.','SOURCES: high_seat_page.py · high_seat_card.py · active-set W2e tests · UI SCHEMATIC',
'''Source and regression definitions support expanded cards with Status, Next steps, Deliverables and Plan, current dialogue, actionable decision context and file drop into the associated cockpit folder. The activity signal is bound to the specific effort's cockpit, not only the parent project. Color semantics: green working, orange needs you, red problem, gray idle. These are descriptions of implemented UI behavior; this deck did not submit a live message, upload a file or retest all authenticated clicks. A September 30 journal records corrected effort links, but authenticated destination/reopen checks remained a separate gate. Show a real known-good card in the demo.''')
rect(s,.8,2.4,8.2,3.75,WHITE,radius=True)
txt(s,1.05,2.65,7.5,.42,'Workshop preparation  ·  needs your input',22,TEAL,True)
rect(s,1.05,3.3,7.65,.88,PALE);txt(s,1.25,3.4,7.25,.73,'Decision: use the beginner example?\nRecommendation + review link stay with the question.',17)
txt(s,1.05,4.4,7.55,.5,'Status   |   Next steps   |   Deliverables   |   Plan',18,TEAL,True)
rect(s,1.05,5.14,7.65,.65,BG);txt(s,1.25,5.28,7.15,.35,'Continue the conversation here…       Drop a source file',17,GRAY)
for i,(a,c) in enumerate([('Working','31836A'),('Needs you','C7892C'),('Problem',RED),('Idle','82909A')]):
    rectshape(s,MSO_SHAPE.OVAL,9.6,2.72+i*.75,.2,.2,c);txt(s,10,2.62+i*.75,2.3,.4,a,20)
band(s,'Moving a card changes your attention list. Opening it preserves the effort’s identity.')

s=new('The AI cockpit keeps the work together','06 · Project and effort workspace','Plan, conversation and outputs remain connected as you refine a deliverable.','SOURCES: steward_cockpit/steward_campaign.py · Ecgberht SKILL.md · UI SCHEMATIC',
'''Ecgberht is the underlying project/portfolio steward. Its durable memory includes a narrative Face (ECGBERHT.md), attention Strip (strip.json) and roadmap/events. Colleagues do not need these filenames; they need to understand that state survives a conversation. Current work and review items should remain easy to open, accepted versions stay available, and older/supporting material stays under its owner. Explicit human acceptance is separate from a successful automated check. Navigation and acceptance behavior vary across legacy campaigns, so avoid claiming universal proof. Demo cue: show one current deliverable, the discussion that produced it and its next step.''')
for x,w,title,body in [( .8,3,'PLAN','Goal and milestones\nCurrent step\nDependencies\nNext action'),(4.05,4.65,'STEWARD CONVERSATION','Give context and ask questions.\nRefine the work.\nCommission a specialist.\nReview and decide.'),(8.95,3.55,'DELIVERABLES','Current work + review\nAccepted versions\nSupporting material\nEarlier refinements')]:
    card(s,x,2.55,w,3.2,title,body,TEAL,19)
band(s,'Project = durable home. Effort = a distinct piece of work. Cockpit = its working interface.')

s=new('Foundry skills: organize, investigate, invent, build','07 · Specialist directory · 1 of 2','Each named skill has a purpose; the steward chooses the appropriate depth and commissions it.','SOURCES: Trio skill definitions · Foundry HUMAN.md cards · Ecgberht SKILL.md',
'''These are capability summaries from local skill cards, not a claim that every task launches all of them. Ecgberht owns campaign continuity and commissions specialists without reimplementing them. ResearchPrime researches and validates; Gandalf assesses; Jumper/Superjump explore ideas; Crucible plans; Foreman executes a frozen plan with checks and corrections. Superjump is the structured deeper ideation path; Jumper remains the lean path. Foundry is also the maintenance/improvement process for the skills: ordinary project work does not silently rewrite their engines.''')
rows=[('Ecgberht / Steward','Remember the campaign; organize and route the next work.'),('Research Prime','Investigate a question; grade evidence and challenge findings.'),('Gandalf','Assess an artifact or system; expose weaknesses and options.'),('Jumper / Superjump','Generate ideas; test fit and discard weak directions.'),('Crucible','Turn a goal into a reviewed, executable plan.'),('Foreman','Build from the plan; run checks, repair and track progress.')]
for i,(a,b) in enumerate(rows):
    y=2.35+i*.65;rect(s,.8,y,11.7,.55,WHITE if i%2==0 else PALE);txt(s,1,y+.1,3.5,.35,a,18,TEAL,True);txt(s,4.5,y+.1,7.7,.35,b,17)

s=new('Foundry skills: domain expertise and upkeep','08 · Specialist directory · 2 of 2','Use the specialist that matches the artifact and the kind of checking it needs.','SOURCES: Foundry skills/*/HUMAN.md and SKILL.md · Anchor health tooling',
'''Domain skills vary in implementation: do not imply every entry is the same multi-agent engine. Ramanujan combines mathematical reasoning with checkable certificates/solvers where applicable. Legal Beagle grounds analysis in supplied or retrieved sources; Financial Analyst builds financial models. Literature Review manages source-grounded literature synthesis. Einstein creates accessible study material; Chekhov preserves story continuity and author decisions. Figure Designer produces scientific/technical figures. Tidy-Idy proposes folder organization with reversible handling. Zombie Hunter inspects ownership before acting on orphan processes. Doctor is Anchor's system-health entry point, shown as an adjacent support tool rather than falsely listed as an identical Foundry domain engine. No current clinical/legal/financial advice is being given here.''')
groups=[('REASON + RESEARCH','Ramanujan — mathematics and proof checks\nLiterature Review — source-grounded synthesis\nLegal Beagle — contracts and legal sources\nFinancial Analyst — models and valuations'),('CREATE + TEACH','Einstein — learning and study packs\nChekhov — story memory and continuity\nFigure Designer — scientific figures'),('KEEP THE SYSTEM USABLE','Tidy-Idy — reviewed folder organization\nZombie Hunter — orphan-process control\nDoctor — Anchor health and diagnosis\nFoundry — improve the skills themselves')]
for i,(a,b) in enumerate(groups):card(s,.8+i*4.02,2.5,3.75,3.7,a,b,TEAL,17)

s=new('A programmatic skill is a managed workflow','09 · What happens behind the conversation','Software controls the stages. AI agents fill specialist roles within those stages.','SOURCES: researchPrime/bin/{run-rounds,governor,round}.mjs · Foundry model policy',
'''Distinguish a prompt skill from a programmatic agent workflow. ResearchPrime and other engine-backed Foundry tools include code that routes roles, validates structured inputs, checkpoints results, limits iterations and handles failure. Agents supply research, judgment, synthesis or coding inside that workflow. Not all skills have identical machinery. The orchestrator is the coordinating layer, not a promise of fully autonomous unlimited action. Reviewer roles do not automatically establish provider independence. Code determines the stages; the models' reasoning and external tools determine what evidence/findings emerge. The operator may need to correct artifacts and supply the next round input.''')
card(s,.8,2.45,3.15,3.3,'ORCHESTRATOR','Chooses stages\nRoutes agent roles\nSaves progress\nEnforces limits\nStops when required',TEAL,20)
arrow(s,4.15,3.88,.5,.3)
card(s,4.85,2.45,3.4,3.3,'AGENT SEATS','Researcher / author\nSkeptic / reviewer\nJudge\nSynthesizer',TEAL,21)
arrow(s,8.45,3.88,.5,.3)
card(s,9.15,2.45,3.35,3.3,'CHECKED OUTPUT','Evidence + artifact\nReview findings\nRevisions\nResult + unresolved issues',TEAL,20)
band(s,'Boundaries include scope, permissions, review depth, iteration budget and human decisions.')

s=new('Research Prime: question → evidence → challenge','10 · Actual workflow, simplified','Example question: “Which approach should our team adopt, and what would change that recommendation?”','SOURCES: researchPrime/SKILL.md · bin/plan-gate.mjs · run-rounds.mjs · round.mjs · governor.mjs',
'''This maps the real protocol, not a ResearchPrime run performed for this deck. Step 1 defines the objective, alternatives, falsification criteria and stakes. Step 2 requires the human's frozen-plan decision. Step 3 executes from that plan, with real context isolation only where the host supports and verifies it; otherwise approximated isolation is disclosed. Step 4 grades evidence. The live round driver spawns Skeptic, Contrarian and Analyst review roles. Deterministic tally and stakes governor precede expensive adjudication. Conditional debate handles conflicting independent-origin findings. Judge and Synthesizer then run concurrently when admitted; the judge tests findings against the objective, while synthesis directs improvement. The operator applies fixes and supplies the next round; the diagram does not claim the driver writes all revisions automatically. Completion includes qualified convergence or an honest budget/failure stop, with open issues preserved. engine.mjs is a skeleton and is NOT the canonical live entry point; run-rounds.mjs is.''')
for i,(a,b) in enumerate([('1  FRAME','Question, options, stakes'),('2  HUMAN GATE','Approve the frozen plan'),('3  INVESTIGATE','Execute from the plan'),('4  EVIDENCE','Grade sources + claims')]):
    card(s,.8+i*3.02,2.4,2.8,1.27,a,b,TEAL,16)
    if i<3:arrow(s,3.63+i*3.02,2.93,.17,.22)
line(s,11.25,3.72,11.25,4.07);rectshape(s,MSO_SHAPE.DOWN_ARROW,11.11,3.86,.28,.28,TEAL)
for x,a,b in [(8.9,'5  ADVERSARIAL PANEL','Skeptic · Contrarian · Analyst'),(4.85,'6  TALLY + ADJUDICATE','Relevance, agreement; judge + synthesis; debate if needed'),(.8,'7  REVISE OR REPORT','Fix material blockers and repeat; otherwise report result + limits')]:
    card(s,x,4.18,3.62,1.68,a,b,TEAL,16)
rectshape(s,MSO_SHAPE.LEFT_ARROW,8.52,4.89,.29,.25,TEAL);rectshape(s,MSO_SHAPE.LEFT_ARROW,4.47,4.89,.29,.25,TEAL)
line(s,2.5,5.92,2.5,6.17,GOLD);line(s,2.5,6.17,10.7,6.17,GOLD);line(s,10.7,6.17,10.7,5.95,GOLD)
txt(s,4.0,6.28,6,.3,'Repeat only for material unresolved issues, within the run budget.',13,TEAL,True)

s=new('What drives the reviews—and what you configure','11 · Judgment and control','A reviewer’s agreement is evidence to inspect, not proof that the answer is correct.','SOURCES: researchPrime governor/round code · Anchor settings · Foundry AGENTS.md',
'''Local policy resolves coding_family, review_family and default_cli from Anchor settings, with a mirror for other hosts. The configured providers use subscription CLI transports. Latest/highest supported model policy is separate from task depth. Same-family coding and review is allowed but must be labeled cross_model:false. Do not imply three reviewer personas are three independent model families. In ResearchPrime, findings must trace to the goal, agreement is keyed to claim identity, the tally uses >=2-agree logic, and independence is separately evaluated from origin/family metadata. Medium/high stakes allow the Judge/Synthesizer/conditional debate; low stakes omit those expensive stages. Code distinguishes empty review from a genuine dry round. One-dry-round policy exists in canonical prose, while generic tracker code accepts preregistered thresholds: do not promise a universal fixed number of rounds. Default max-rounds is eight in the documented driver, but budget limits are run controls, not all exposed as GUI sliders.''')
card(s,.8,2.45,3.75,3.65,'QUALITY DRIVERS','Does the finding affect the goal?\nHow strong is its evidence?\nDo reviewers identify the same claim?\nAre their origins independent?',TEAL,19)
card(s,4.8,2.45,3.75,3.65,'ANCHOR SETTINGS','Choose the interactive provider.\nChoose author / coding family.\nChoose review family.\nSeats follow those preferences.',TEAL,19)
card(s,8.8,2.45,3.75,3.65,'WORKFLOW RULES','Stakes determine review depth.\nMaterial issues trigger revision.\nBudgets bound repetition.\nFailures and missing checks stay visible.',GOLD,19)
band(s,'Evidence ladder: observed → corroborated → claimed → unverified → refuted.')

s=new('The earlier best-in-class report: useful, but dated','12 · Read the assessment in context','Its central distinction still helps: strong internal machinery must translate into an easy, reliable experience.','SOURCES: Ecgberht/planning/steward-assessment-2026-08-08/ASSESSMENT.md · newer code and friction records',
'''The closest directly relevant local report located is “Stage 1 — Honest assessment of the steward,” dated August 8, 2026. Also reviewed Skill Foundry/best-in-class-action-plan.md, which is a remediation proposal, not a benchmark certification. The assessment praised ledger/receipt/budget architecture and criticized fragmented status, 60–160-second cold turns and lack of a declared execution pipeline. Its claim of unrivaled standing was not a measured competitive study and is not repeated here as fact. Newer high_seat_page.py explicitly renders the portfolio from records without a model call; current expanded cards and shared status composition address parts of the old UI critique. No fresh latency benchmark was run. September 30/October 1 friction documents still show acknowledged unfinished work stranded after interruption. Thus the old report should guide verification, not be used as a current score.''')
card(s,.8,2.45,3.75,3.55,'WHAT IT PRAISED','Durable campaign records\nReceipts and bounded work\nSpecialist commissioning\nWhole-campaign context',TEAL,20)
card(s,4.8,2.45,3.75,3.55,'WHAT IT CRITICIZED','Fragmented status\nSlow cold conversation turns\nWeak flow between stages\nUnclear “what happens next”',RED,20)
card(s,8.8,2.45,3.75,3.55,'WHAT HAS MOVED','Record-driven home rendering\nExpandable initiative cards\nShared cockpit status logic\nRecovery gaps still documented',TEAL,19)
band(s,'Assessment dated 8 August 2026. Current comparison below is qualitative, not a competitive ranking.')

s=new('How close is Anchor to an excellent assistant?','13 · Evidence-based scorecard','Strong architectural coverage; reliability and everyday usability still need end-to-end proof.','BENCHMARK: official LangGraph + Anthropic engineering guidance · local evidence in speaker notes',
'''Benchmark is a synthesis, not a published ranking. Official sources: https://docs.langchain.com/oss/python/langgraph/persistence ; https://docs.langchain.com/oss/python/concepts/memory ; https://www.anthropic.com/engineering/building-effective-agents ; https://docs.langchain.com/oss/python/langgraph/streaming ; https://docs.langchain.com/oss/python/langgraph/interrupts ; https://www.anthropic.com/engineering/demystifying-evals-for-ai-agents . Anchor evidence: high_seat_page.py, campaign_page.py, steward_campaign.py, ResearchPrime governor/round code, Foundry model policy, September 30–October 1 interruption friction. “Present in source” means code/records inspected, not all-device runtime certification. Recovery is partial because native history can be retained while acknowledged unfinished work still strands. Visibility is partial because current features and recent corrections exist, with remaining live consistency verification. Integrations are partial: email/calendar explicitly not connected in the current renderer. Quality measurement is unproven because this review ran no representative cross-product task suite or latency study. Do not invent a percent complete.''')
rows=[('Campaign memory','Present in source','Goals, decisions and effort records persist.'),('Specialist orchestration','Present in source','Programmed stages, reviews and model routing.'),('Human control','Present in source','Plan decisions, scoped work and review boundaries.'),('Recovery + status consistency','Partial','Recent interrupted-work and stale-status friction.'),('Connected assistant services','Partial','Multi-device access; email/calendar not connected.'),('Measured real-world quality','Needs evidence','Task success, recovery, latency and mobile benchmarks.')]
for i,(a,b,c) in enumerate(rows):
    y=2.35+i*.65;rect(s,.8,y,11.72,.56,WHITE if i%2==0 else PALE)
    txt(s,.95,y+.1,3.2,.38,a,16,NAVY,True);txt(s,4.22,y+.1,2.35,.38,b,15,TEAL if b=='Present in source' else RED,True);txt(s,6.7,y+.1,5.6,.38,c,15)
band(s,'A promising personal work system. “Best in class” remains a claim to demonstrate with outcomes.')

s=new('Show one complete journey in the live demo','14 · Demonstration and next priorities','Follow one initiative from its goal to a useful output—and make the remaining gaps concrete.','RECOMMENDED DEMO · fictional example content · priorities derived from source review, not a build plan',
'''Suggested 5–7 minute demo: 1) open the Cave of Wonders on one device; 2) show the Campaign goal and a linked initiative; 3) expand its Active card and show the exact effort; 4) open current output, pending decision and conversation; 5) show author/reviewer preferences and a saved ResearchPrime result, not a long new run. If useful, switch to another authorized device to show the same workspace. Before claiming dependable recovery, perform a separately authorized safe acceptance exercise on a non-disruptive test job; do not restart the shared service during this presentation task. Highest priorities: durable acknowledged-work recovery; consistent current status and exact effort navigation; representative task, latency and mobile checks; real email/calendar integration and accounting completeness. The deck is Initial mockup v1 for the user's review; no approval of software changes or publication is implied.''')
card(s,.8,2.45,6.05,3.65,'THE FIVE-MINUTE TOUR','1  Open the Cave of Wonders.\n2  Show one goal and its initiative.\n3  Expand the exact effort’s card.\n4  Open its dialogue and deliverable.\n5  Show a specialist’s reviewed output.',TEAL,21)
card(s,7.15,2.45,5.35,3.65,'WHAT WOULD MAKE IT STRONGER','Reliable recovery of unfinished work\nConsistent status and exact effort links\nMeasured task and mobile performance\nConnected email/calendar and complete usage accounting',GOLD,20)
band(s,'The value to show: less re-briefing, clearer priorities, and work that stays connected to its purpose.')

R.core_properties.title='Anchor — Colleague Orientation';R.core_properties.subject='Initial mockup v1';R.core_properties.author='John';R.core_properties.comments='Source-based orientation. Illustrative UI schematics. Not an acceptance certificate.'
R.save(OUT/'Anchor-Orientation-Mockup-v1.pptx')
(OUT/'slide-notes.json').write_text(json.dumps(slides,indent=2,ensure_ascii=False),encoding='utf-8')
body=''.join(f'<section><h2>{i+1}. {html.escape(a["title"])}</h2><p>{html.escape(a["notes"])}</p><small>{html.escape(a["source"])}</small></section>' for i,a in enumerate(slides))
(OUT/'Presenter-notes-and-sources.html').write_text('<!doctype html><meta charset="utf-8"><title>Anchor orientation — notes and sources</title><style>body{font:18px/1.6 system-ui;max-width:950px;margin:50px auto;padding:0 25px;color:#152d3a}section{margin:40px 0;border-top:1px solid #ccc}small{color:#556671}</style><h1>Anchor orientation · Initial mockup v1</h1><p>15-slide colleague presentation. Schematic panels use fictional examples. Source review: 1 October 2026. Full notes are also embedded in the PowerPoint.</p>'+body,encoding='utf-8')
print(f'Created {len(R.slides)} slides: {OUT}')
