"""Make the reading copy of a Markdown document in seconds: the same-version PDF (and, asked for, HTML) beside it.

  python C:/dev/Anchor/tools/reading_copy.py <document.md> [--html] [--out <file.pdf>]

John 2026-10-12: agents spent 16-19 minutes writing a PDF renderer for every plan and outline. This is the one
renderer: Anchor's own Markdown page (report_viewer, math included), printed by the Edge already on this machine.
Prints the written path(s); exit 1 with the reason when it could not."""
import argparse
import base64
import html as htmllib
import urllib.parse
import mimetypes
import re
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))
import report_viewer  # noqa: E402

EDGE = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"]

PRINT_CSS = ("<style>body{background:#fff !important;color:#111 !important;max-width:none !important;"
             "padding:0 !important;font-size:11pt} h1,h2,h3,h4,a{color:#1a4fb5 !important}"
             " code,pre,th{background:#f2f4f8 !important;color:#111 !important} th,td{border-color:#bbb !important}"
             " @page{margin:18mm 16mm}</style>\n</head>")


def page(md_path, for_print):
    text = md_path.read_text(encoding="utf-8", errors="replace")
    title = next((ln.lstrip("# ").strip() for ln in text.splitlines() if ln.startswith("#")), md_path.stem)
    html = report_viewer.reader_html(text, title=title, katex_url_prefix=report_viewer.VENDOR_KATEX_DIR.as_uri())
    # review r7: the reader renders ![alt](src) as "!" + a link; a reading copy shows the figure itself
    html = re.sub(r'!<a href="([^"]+)" rel="noopener noreferrer">([^<]*)</a>',
                  r'<img src="\1" alt="\2" style="max-width:100%">', html)
    if not for_print:
        # review r8: Anchor opens the HTML copy through a file route, where a relative figure path breaks - local
        # figures from the document's own folder ride inside the copy
        folder = md_path.resolve().parent

        def embed(m):
            src = urllib.parse.unquote(htmllib.unescape(m.group(1)))      # review r9: "Figure 1.png" arrives escaped
            try:
                f = (folder / src).resolve()
                f.relative_to(folder)
                kind = mimetypes.guess_type(f.name)[0] or ""
                if kind.startswith("image/") and f.is_file() and f.stat().st_size <= 8 * 1024 * 1024:
                    return '<img src="data:%s;base64,%s"' % (kind, base64.b64encode(f.read_bytes()).decode("ascii"))
            except (OSError, ValueError):
                pass
            return m.group(0)
        return re.sub(r'<img src="([^":]+)"', embed, html)
    # review r4: the print page is a temp file - its relative links resolve against the document's own folder
    base = '<base href="%s/">\n' % md_path.resolve().parent.as_uri()
    return html.replace("<head>\n", "<head>\n" + base, 1).replace("</head>", PRINT_CSS, 1)


def edge():
    for p in EDGE + [shutil.which("msedge") or ""]:
        if p and Path(p).is_file():
            return p
    return None


def _write_docx(md: Path, docx_out: Path) -> bool:
    try:
        import importlib.util
        tool_py = Path(os.environ.get("ANCHOR_FOUNDRY_TOOLS", "C:/dev/foundry-tools")) / "docx-redline" / "tool.py"
        if tool_py.is_file():
            spec = importlib.util.spec_from_file_location("docx_redline_tool", str(tool_py))
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)
            mod.markdown_to_docx(md, docx_out)
            return docx_out.is_file() and docx_out.stat().st_size > 0
        from docx import Document
        doc = Document()
        for raw in md.read_text(encoding="utf-8", errors="replace").splitlines():
            s = raw.strip()
            if not s:
                continue
            if s in ("---", "***", "* * *", "___"):
                p = doc.add_paragraph("* * *")
                p.alignment = 1
            elif s.startswith("#"):
                lvl = min(4, len(s) - len(s.lstrip("#")))
                doc.add_heading(s[lvl:].strip(), level=lvl)
            else:
                doc.add_paragraph(s)
        docx_out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(docx_out))
        return docx_out.is_file() and docx_out.stat().st_size > 0
    except Exception as e:
        print("no DOCX was made: %s" % e)
    return False


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("document")
    ap.add_argument("--html", nargs="?", const="", default=None,
                    help="also write the same-version HTML (beside it, or at the path given)")
    ap.add_argument("--docx", nargs="?", const="", default=None,
                    help="also write the same-version DOCX (beside it, or at the path given)")
    ap.add_argument("--out", help="the PDF path (default: the document's name with .pdf)")
    ap.add_argument("--no-pdf", action="store_true", help="only the HTML and/or DOCX")
    a = ap.parse_args(argv)
    md = Path(a.document).resolve()
    if not md.is_file():
        print("no such document: %s" % md)
        return 1
    html_out = None if a.html is None else (Path(a.html).resolve() if a.html else md.with_suffix(".html"))
    docx_out = None if a.docx is None else (Path(a.docx).resolve() if a.docx else md.with_suffix(".docx"))
    if a.no_pdf:
        if html_out is None and docx_out is None:
            print("--no-pdf needs --html or --docx")
            return 1
        if html_out is not None:
            html_out.write_text(page(md, False).replace(report_viewer.VENDOR_KATEX_DIR.as_uri(), "/vendor/katex"),
                                encoding="utf-8")
            print(html_out)
        if docx_out is not None:
            if not _write_docx(md, docx_out):
                return 1
            print(docx_out)
        return 0
    pdf = Path(a.out).resolve() if a.out else md.with_suffix(".pdf")
    browser = edge()
    if not browser:
        print("Edge was not found, so no PDF was made")
        return 1
    with tempfile.TemporaryDirectory(prefix="reading-copy-") as tmp:
        src = Path(tmp, "page.html")
        src.write_text(page(md, True), encoding="utf-8")
        profile = Path(tmp, "profile")
        fresh = Path(tmp, "out.pdf")              # review r2: a failed render never passes off an older PDF as new
        cmd = [browser, "--headless=new", "--disable-gpu", "--no-first-run", "--no-pdf-header-footer",
               "--user-data-dir=" + str(profile), "--virtual-time-budget=3000", "--print-to-pdf=" + str(fresh),
               src.as_uri()]
        try:
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=90,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except subprocess.TimeoutExpired:
            print("Edge did not finish the PDF within 90 s")
            return 1
        if not fresh.is_file() or fresh.stat().st_size < 500 or not fresh.read_bytes().startswith(b"%PDF"):
            print("no PDF was made (Edge exit %s): %s" % (r.returncode, (r.stderr or "").strip()[-300:]))
            return 1
        pdf.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(fresh, pdf)
    print(pdf)
    if html_out is not None:
        html_out.write_text(page(md, False).replace(report_viewer.VENDOR_KATEX_DIR.as_uri(), "/vendor/katex"),
                            encoding="utf-8")
        print(html_out)
    if docx_out is not None:
        if not _write_docx(md, docx_out):
            return 1
        print(docx_out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
